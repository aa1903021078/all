#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""对 success.docx 做最终校验: 字数 / 重复段落 / 图片 / 标题 / 表格。"""
import re
from docx import Document
from docx.oxml.ns import qn

DOCX = "/Users/weibaiwang/Desktop/AAA/work/all/success.docx"
doc = Document(DOCX)

# 1) 字数统计
cjk = 0
word_like = 0
all_text = []
for p in doc.paragraphs:
    all_text.append(p.text)
for t in doc.tables:
    for row in t.rows:
        for c in row.cells:
            all_text.append(c.text)
joined = "\n".join(all_text)
cjk = len(re.findall(r'[\u4e00-\u9fff]', joined))
# Word 风格: 中文字符(含标点近似) + 英文单词 + 数字串
en_tokens = len(re.findall(r'[A-Za-z]+', joined))
num_tokens = len(re.findall(r'\d+', joined))
cjk_punct = len(re.findall(r'[\u3000-\u303f\uff00-\uffef]', joined))
word_style = cjk + cjk_punct + en_tokens + num_tokens
print("== 字数 ==")
print("  纯中文字符(CJK):", cjk)
print("  Word近似字数(中文+中文标点+英文词+数字):", word_style)

# 2) 连续重复段落检测
print("== 连续重复段落检测 ==")
dup = 0
prev = None
for p in doc.paragraphs:
    txt = p.text.strip()
    if txt and txt == prev and len(txt) > 15:
        dup += 1
        print("  [重复]", txt[:40], "...")
    prev = txt
if dup == 0:
    print("  未发现连续重复正文段落 ✓")

# 3) 图片数量
img_count = 0
for rel in doc.part.rels.values():
    if "image" in rel.reltype:
        img_count += 1
print("== 图片 ==")
print("  内嵌图片数量:", img_count)

# 4) 图题编号
caps = re.findall(r'图(\d+)-(\d+)', joined)
tabs = re.findall(r'表(\d+)-(\d+)', joined)
print("  图题引用/标注(图X-Y)出现次数:", len(caps))
print("  表题引用/标注(表X-Y)出现次数:", len(tabs))

# 5) 标题(outline 级别)
print("== 标题大纲级别 ==")
h1 = h2 = h3 = 0
for p in doc.paragraphs:
    pPr = p._p.find(qn('w:pPr'))
    if pPr is None:
        continue
    lvl = pPr.find(qn('w:outlineLvl'))
    if lvl is not None:
        v = lvl.get(qn('w:val'))
        if v == '0':
            h1 += 1
        elif v == '1':
            h2 += 1
        elif v == '2':
            h3 += 1
print("  一级(outline0):", h1, " 二级(outline1):", h2, " 三级(outline2):", h3)

# 6) 表格
print("== 表格 ==")
print("  表格数量:", len(doc.tables))

# 7) 页面设置
sec = doc.sections[0]
print("== 页面 ==")
print("  纸张(EMU) 宽:", sec.page_width, " 高:", sec.page_height)
print("  段落总数:", len(doc.paragraphs))
