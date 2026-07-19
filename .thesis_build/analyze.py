#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""分析模版.docx：批注、样式、结构"""
import sys, re
from docx import Document

TPL = "/Users/weibaiwang/Desktop/AAA/work/all/模版.docx"

def analyze_comments():
    """解析 word/comments.xml 中的批注"""
    import zipfile
    from lxml import etree
    ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    with zipfile.ZipFile(TPL) as z:
        data = z.read('word/comments.xml')
    root = etree.fromstring(data)
    print("=" * 60)
    print("批注 (COMMENTS) 列表：")
    print("=" * 60)
    for c in root.findall('w:comment', ns):
        cid = c.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}id')
        author = c.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}author')
        texts = c.findall('.//w:t', ns)
        content = ''.join(t.text or '' for t in texts)
        print(f"[#{cid}] {content}")

def analyze_styles():
    doc = Document(TPL)
    print("\n" + "=" * 60)
    print("样式 (STYLES) 列表：")
    print("=" * 60)
    for s in doc.styles:
        try:
            print(f"- {s.name}  | type={s.type}")
        except Exception:
            pass

def analyze_structure():
    doc = Document(TPL)
    print("\n" + "=" * 60)
    print("文档结构 (段落样式 + 文本预览)：")
    print("=" * 60)
    for i, p in enumerate(doc.paragraphs):
        txt = (p.text or '').strip()
        st = p.style.name if p.style else '?'
        if txt or st.lower().startswith('heading') or 'title' in st.lower() or '标题' in st:
            preview = txt[:50]
            print(f"[{i}] <{st}> {preview}")

def analyze_sections():
    doc = Document(TPL)
    print("\n" + "=" * 60)
    print("页面设置 (SECTIONS)：")
    print("=" * 60)
    for i, sec in enumerate(doc.sections):
        print(f"Section {i}: page={sec.page_width}x{sec.page_height} "
              f"margins L={sec.left_margin} R={sec.right_margin} T={sec.top_margin} B={sec.bottom_margin}")

if __name__ == '__main__':
    analyze_comments()
    analyze_styles()
    analyze_structure()
    analyze_sections()
