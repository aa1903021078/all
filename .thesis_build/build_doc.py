#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""构建 success.docx: 封面 + 中英文摘要 + 目录 + 八章正文 + 参考文献 + 致谢。"""
import sys, re
sys.path.insert(0, "/Users/weibaiwang/Desktop/AAA/work/all/.thesis_build")
from docx_util import *
from docx_util import _cell_border
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from content_a import BLOCKS_1_4
from content_b import BLOCKS_5_8

OUT_DOCX = "/Users/weibaiwang/Desktop/AAA/work/all/success.docx"
LOGO = f"{IMG}/school_logo.jpeg"

TITLE = "基于 Spring Boot 与 Vue 的美食探店与菜谱分享平台的设计与实现"

ABS_CN = [
    "随着本地生活服务与内容社区的深度融合,餐饮消费正从单一的到店就餐向社交化探店与家庭化烹饪多场景延伸,而传统平台大多将探店点评与菜谱分享割裂开来,难以形成完整的用户体验闭环。针对这一问题,本文设计并实现了一套融合探店与菜谱的一体化平台“觅食记”,以提升用户体验并为餐饮商家提供新的经营渠道。",
    "系统采用 Spring Boot 与 Vue 前后端分离架构,后端基于 MyBatis-Plus 访问 MySQL 数据库,采用 JWT 无状态鉴权并结合基于角色的访问控制保障安全,通过 WebSocket 实现用户与商家的实时聊天;前端基于 Vue 3、Element Plus 与 ECharts 构建。系统面向普通用户、商家与管理员三类角色,实现了首页推荐、美食地图、探店笔记、在线预约、点亮打卡、菜谱分享、食材采购清单、实时聊天以及店铺审核、内容审核、运营大屏等功能模块。",
    "经功能测试与性能测试验证,系统各功能模块运行正常,权限控制有效,核心接口在一定并发下响应迅速、运行稳定,达到了预期的设计目标,具有良好的实用价值与推广前景。",
]
KW_CN = "美食探店；菜谱分享；Spring Boot；Vue;前后端分离"

ABS_EN = [
    "With the deep integration of local life services and content communities, dining consumption is extending from a single dine-in scenario to socialized restaurant exploration and home cooking. However, traditional platforms usually separate restaurant reviews from recipe sharing, making it difficult to form a complete experience loop. To address this problem, this paper designs and implements an integrated platform named MiShiJi that combines restaurant exploration and recipe sharing.",
    "The system adopts a front-end and back-end separation architecture based on Spring Boot and Vue. The back end accesses a MySQL database through MyBatis-Plus, uses stateless JWT authentication together with role-based access control to ensure security, and realizes real-time chat between users and merchants through WebSocket. The front end is built with Vue 3, Element Plus and ECharts. Serving ordinary users, merchants and administrators, the system implements modules such as home recommendation, food map, exploration notes, online reservation, shop check-in, recipe sharing, ingredient shopping list, real-time chat, shop auditing, content auditing and an operation dashboard.",
    "Functional and performance tests show that all modules of the system work correctly, the access control is effective, and the core interfaces respond quickly and run stably under a certain concurrency, which meets the expected design goals and shows good practical value and promotion prospects.",
]
KW_EN = "Food Exploration; Recipe Sharing; Spring Boot; Vue; Front-end and Back-end Separation"


def _center(doc, before=0, after=0, line=1.0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_spacing(p, before=before, after=after, line=line)
    return p


def build_cover(doc):
    for _ in range(2):
        add_blank(doc, 16)
    # 校名
    for line_txt in ["新疆大学软件学院", "小学期实训论文"]:
        p = _center(doc, after=6, line=1.2)
        r = p.add_run(line_txt)
        style_run(r, cn=HEI, en=TNR, size=26, bold=True)
        set_char_spacing(r, 2)
    add_blank(doc, 12)
    # 校徽
    try:
        p = _center(doc, before=6, after=6)
        p.add_run().add_picture(LOGO, width=Cm(3.2))
    except Exception as e:
        print("logo skip:", e)
    for _ in range(2):
        add_blank(doc, 12)
    # 信息表(标签 + 下划线值)
    info = [("论文题目", TITLE),
            ("学生姓名", ""),
            ("学    号", ""),
            ("所属院系", "软件学院"),
            ("专    业", "软件工程"),
            ("班    级", ""),
            ("指导老师", ""),
            ("日    期", "2026 年      月      日")]
    table = doc.add_table(rows=len(info), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (label, val) in enumerate(info):
        lc = table.rows[i].cells[0]
        vc = table.rows[i].cells[1]
        lc.width = Cm(3.2); vc.width = Cm(9.0)
        # 标签
        lp = lc.paragraphs[0]; lp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        set_spacing(lp, before=4, after=4, line=1.3)
        lr = lp.add_run(label + " ："); style_run(lr, cn=SONG, en=TNR, size=SZ["四号"], bold=True)
        # 值(带下边框下划线)
        vp = vc.paragraphs[0]; vp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_spacing(vp, before=4, after=4, line=1.3)
        sz = SZ["小四"] if label == "论文题目" else SZ["四号"]
        vr = vp.add_run(val if val else "　　　　　　"); style_run(vr, cn=SONG, en=TNR, size=sz)
        _cell_border(vc, 'bottom', 6)
    # 去掉信息表其他边框(默认无, Normal Table)
    page_break_after_paragraph(doc)


def page_break_after_paragraph(doc):
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


def build_abstract_cn(doc):
    # 摘  要 标题
    p = _center(doc, before=12, after=14, line=1.0)
    r = p.add_run("摘    要"); style_run(r, cn=HEI, en=TNR, size=SZ["三号"], bold=True)
    set_outline_level(p, 0)
    for para in ABS_CN:
        add_body(doc, para)
    # 关键词
    kp = doc.add_paragraph(); set_spacing(kp, before=12, after=0, line=1.5)
    set_first_line_indent_chars(kp, 2)
    r1 = kp.add_run("关键词："); style_run(r1, cn=SONG, en=TNR, size=SZ["小四"], bold=True)
    r2 = kp.add_run(KW_CN); style_run(r2, cn=SONG, en=TNR, size=SZ["小四"])
    page_break_after_paragraph(doc)


def build_abstract_en(doc):
    p = _center(doc, before=12, after=14, line=1.0)
    r = p.add_run("ABSTRACT"); style_run(r, cn=TNR, en=TNR, size=SZ["三号"], bold=True)
    set_outline_level(p, 0)
    for para in ABS_EN:
        bp = doc.add_paragraph(); bp.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        set_spacing(bp, before=0, after=0, line=1.5)
        set_first_line_indent_chars(bp, 2)
        rr = bp.add_run(para); style_run(rr, cn=TNR, en=TNR, size=SZ["小四"])
    kp = doc.add_paragraph(); set_spacing(kp, before=12, after=0, line=1.5)
    set_first_line_indent_chars(kp, 2)
    r1 = kp.add_run("KEY WORDS: "); style_run(r1, cn=TNR, en=TNR, size=SZ["小四"], bold=True)
    r2 = kp.add_run(KW_EN); style_run(r2, cn=TNR, en=TNR, size=SZ["小四"])
    page_break_after_paragraph(doc)


def build_toc_section(doc):
    p = _center(doc, before=12, after=14, line=1.0)
    r = p.add_run("目    录"); style_run(r, cn=HEI, en=TNR, size=SZ["三号"], bold=True)
    add_toc(doc)
    page_break_after_paragraph(doc)


def add_reference(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    set_spacing(p, before=0, after=0, line=1.5)
    # 悬挂缩进
    pPr = p._p.get_or_add_pPr()
    ind = OxmlElement('w:ind')
    ind.set(qn('w:left'), "480"); ind.set(qn('w:hanging'), "480")
    pPr.append(ind)
    r = p.add_run(text); style_run(r, cn=SONG, en=TNR, size=SZ["小四"])
    return p


def render_blocks(doc, blocks):
    for b in blocks:
        kind = b[0]
        if kind == 'h1':
            add_h1(doc, b[1], page_break=True)
        elif kind == 'h2':
            add_h2(doc, b[1])
        elif kind == 'h3':
            add_h3(doc, b[1])
        elif kind == 'p':
            add_body(doc, b[1])
        elif kind == 'fig':
            add_figure(doc, b[1], b[2], b[3] if len(b) > 3 else 13.0)
        elif kind == 'table':
            add_three_line_table(doc, b[1], b[2], b[3], b[4] if len(b) > 4 else None)
        elif kind == 'code':
            add_code_block(doc, b[2], caption=b[1])
        elif kind == 'refh':
            add_h1(doc, b[1], page_break=True)
        elif kind == 'ref':
            add_reference(doc, b[1])


def count_words(doc):
    cn = 0
    for p in doc.paragraphs:
        cn += len(re.findall(r'[\u4e00-\u9fff]', p.text))
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                cn += len(re.findall(r'[\u4e00-\u9fff]', c.text))
    return cn


def main():
    doc = Document()
    section = doc.sections[0]
    setup_page(section)
    set_header_footer(section, "软件工程专业实训")
    # 正文默认样式
    normal = doc.styles['Normal']
    normal.font.name = TNR
    normal.font.size = Pt(SZ["小四"])
    normal._element.rPr.rFonts.set(qn('w:eastAsia'), SONG)

    build_cover(doc)
    build_abstract_cn(doc)
    build_abstract_en(doc)
    build_toc_section(doc)
    render_blocks(doc, BLOCKS_1_4)
    render_blocks(doc, BLOCKS_5_8)

    enable_update_fields(doc)
    doc.save(OUT_DOCX)
    wc = count_words(doc)
    print("saved:", OUT_DOCX)
    print("中文字符数(不含代码英文):", wc)
    print("段落数:", len(doc.paragraphs), " 表格数:", len(doc.tables))


if __name__ == '__main__':
    main()
